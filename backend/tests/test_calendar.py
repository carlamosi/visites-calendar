import pytest
from fastapi.testclient import TestClient
from backend.main import app
import io
import openpyxl

client = TestClient(app)

def create_mock_excel() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EstudioMock"
    
    # Headers
    ws.cell(row=2, column=3, value="101")
    ws.cell(row=2, column=4, value="102")
    
    # Visit names
    ws.cell(row=4, column=1, value="Basal")
    ws.cell(row=5, column=1, value="Week 1")
    
    # Dates
    from datetime import date
    ws.cell(row=4, column=3, value=date(2026, 1, 1))
    ws.cell(row=5, column=3, value=date(2026, 1, 8))
    ws.cell(row=4, column=4, value=date(2026, 1, 2))
    
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def test_health_check():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}

def test_generate_calendar_valid_file():
    excel_bytes = create_mock_excel()
    response = client.post(
        "/api/generate",
        files={"file": ("mock.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/calendar; charset=utf-8"
    assert "BEGIN:VCALENDAR" in response.text
    assert "SUMMARY:Basal - EstudioMock - 101" in response.text
    assert "SUMMARY:Week 1 - EstudioMock - 101" in response.text
    assert "SUMMARY:Basal - EstudioMock - 102" in response.text

def test_generate_calendar_invalid_extension():
    response = client.post(
        "/api/generate",
        files={"file": ("mock.txt", b"dummy content", "text/plain")}
    )
    assert response.status_code == 400
    assert "Invalid file type" in response.text
