import io
from datetime import datetime, date
from openpyxl import load_workbook
from typing import List

from backend.utils.config import settings
from backend.utils.exceptions import InvalidFileException, EmptyWorkbookException, ParsingException
from backend.models.data_models import CalendarEvent

class ExcelParserService:
    def parse_2d_array(self, rows: List[List[str]], estudio: str) -> List[CalendarEvent]:
        events = []
        has_valid_data = False

        if len(rows) < settings.FILA_PACIENTES:
            raise EmptyWorkbookException(f"The table doesn't have enough rows. Got {len(rows)}, need {settings.FILA_PACIENTES}.")

        fila_pacientes_idx = settings.FILA_PACIENTES - 1
        paciente_row = rows[fila_pacientes_idx]
        
        col_inicio_idx = settings.COL_INICIO_PACIENTES - 1
        columnas_pacientes = []
        
        # DEBUG: log what we see
        import sys
        print(f"[DEBUG] Total rows: {len(rows)}", file=sys.stderr)
        print(f"[DEBUG] fila_pacientes_idx: {fila_pacientes_idx}", file=sys.stderr)
        print(f"[DEBUG] paciente_row (len={len(paciente_row)}): {paciente_row}", file=sys.stderr)
        print(f"[DEBUG] col_inicio_idx: {col_inicio_idx}", file=sys.stderr)
        for i, row in enumerate(rows[:5]):
            print(f"[DEBUG] rows[{i}]: {row}", file=sys.stderr)
        
        for col_idx in range(col_inicio_idx, len(paciente_row)):
            valor = paciente_row[col_idx]
            if valor is None or str(valor).strip() == "":
                break
            columnas_pacientes.append((col_idx, str(valor).strip()))
        
        if not columnas_pacientes:
            raise EmptyWorkbookException(f"No patient columns found in the table. paciente_row={paciente_row}, col_inicio_idx={col_inicio_idx}")


        fila_inicio_visitas_idx = settings.FILA_INICIO_VISITAS - 1
        
        def parse_date(date_str: str) -> date:
            date_str = date_str.strip()
            # Try some common formats
            for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
                try:
                    return datetime.strptime(date_str, fmt).date()
                except ValueError:
                    pass
            # if we get here, it might be something else
            return None

        for fila_idx in range(fila_inicio_visitas_idx, len(rows)):
            row = rows[fila_idx]
            if not row or len(row) == 0:
                continue
                
            nombre_visita = row[0]
            if nombre_visita is None or str(nombre_visita).strip() == "":
                continue
                
            for col_idx, num_paciente in columnas_pacientes:
                if col_idx < len(row):
                    fecha_valor = row[col_idx]
                    
                    if fecha_valor is None or str(fecha_valor).strip() == "":
                        continue
                        
                    parsed_date = parse_date(str(fecha_valor))
                    if not parsed_date:
                        continue
                        
                    has_valid_data = True
                    
                    events.append(
                        CalendarEvent(
                            estudio=estudio,
                            num_paciente=num_paciente,
                            fila=fila_idx + 1,
                            nombre_visita=str(nombre_visita).strip(),
                            fecha=parsed_date
                        )
                    )

        if not has_valid_data and not events:
            raise EmptyWorkbookException("The table contains no valid patient visits.")

        return events

    def parse_excel_bytes(self, file_bytes: bytes) -> List[CalendarEvent]:
        """
        Parses the Excel file bytes into a list of CalendarEvent objects
        matching the original business logic, optimized for performance and memory.
        """
        try:
            # Using read_only=True is critical for fast loading and low memory footprint
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as e:
            raise InvalidFileException(f"Failed to read Excel file: {str(e)}")

        events = []
        has_valid_data = False

        try:
            for ws in wb.worksheets:
                estudio = ws.title  # nombre del ensayo/estudio
                
                # Fetch only the values to keep it fast
                rows = list(ws.iter_rows(values_only=True))
                
                if len(rows) < settings.FILA_PACIENTES:
                    continue
                
                # 0-indexed row for patients
                fila_pacientes_idx = settings.FILA_PACIENTES - 1
                paciente_row = rows[fila_pacientes_idx]
                
                # Detect patient columns
                col_inicio_idx = settings.COL_INICIO_PACIENTES - 1
                columnas_pacientes = []
                
                for col_idx in range(col_inicio_idx, len(paciente_row)):
                    valor = paciente_row[col_idx]
                    if valor is None or str(valor).strip() == "":
                        break
                    columnas_pacientes.append((col_idx, str(valor).strip()))
                
                if not columnas_pacientes:
                    continue
                
                fila_inicio_visitas_idx = settings.FILA_INICIO_VISITAS - 1
                
                for fila_idx in range(fila_inicio_visitas_idx, len(rows)):
                    row = rows[fila_idx]
                    if not row or len(row) == 0:
                        continue
                        
                    nombre_visita = row[0] # column A
                    if nombre_visita is None:
                        continue
                        
                    for col_idx, num_paciente in columnas_pacientes:
                        if col_idx < len(row):
                            fecha_valor = row[col_idx]
                            
                            if fecha_valor is None:
                                continue
                            if not isinstance(fecha_valor, (datetime, date)):
                                continue
                                
                            if isinstance(fecha_valor, datetime):
                                fecha_valor = fecha_valor.date()
                                
                            has_valid_data = True
                            
                            events.append(
                                CalendarEvent(
                                    estudio=estudio,
                                    num_paciente=num_paciente,
                                    fila=fila_idx + 1,
                                    nombre_visita=str(nombre_visita).strip(),
                                    fecha=fecha_valor
                                )
                            )
                            
        except Exception as e:
            raise ParsingException(f"Error parsing worksheet data: {str(e)}")
        finally:
            # Ensure workbook is closed to free resources (important for read_only)
            wb.close()

        if not has_valid_data and not events:
            raise EmptyWorkbookException("The workbook contains no valid patient visits.")

        return events
